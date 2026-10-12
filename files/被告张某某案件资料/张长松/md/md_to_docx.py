#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 Markdown 质证意见转换为 WPS/Word 可打开的 .docx 文档。
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


def set_cell_border(cell, **kwargs):
    """设置表格单元格边框。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = parse_xml(r'<w:{} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'.format(edge))
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))


def set_run_font(run, font_name='仿宋', size=11, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(doc, text, style=None, bold=False, italic=False, size=10.5, align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    return p


def parse_inline_format(paragraph, text):
    """解析 Markdown 行内格式：加粗 **text** 和倾斜 *text*。"""
    # 先处理加粗
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run)
            run.italic = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_markdown_paragraph(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    parse_inline_format(p, text)
    return p


def add_table_from_markdown(doc, lines, start_idx):
    """解析 Markdown 表格，返回（下一行索引，是否处理了表格）。"""
    header_line = lines[start_idx].strip()
    if not header_line.startswith('|'):
        return start_idx, False

    # 收集连续表格行
    table_lines = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        table_lines.append(lines[i].strip())
        i += 1

    if len(table_lines) < 2:
        return start_idx, False

    # 忽略分隔行（如 |---|---|）
    data_lines = [l for l in table_lines if not re.match(r'^\|[-:\s|]+\|$', l.strip())]
    if len(data_lines) < 1:
        return start_idx, False

    rows = []
    for line in data_lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)

    if not rows:
        return start_idx, False

    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'

    for row_idx, row_cells in enumerate(rows):
        row = table.rows[row_idx]
        for col_idx in range(num_cols):
            cell = row.cells[col_idx]
            cell.text = row_cells[col_idx] if col_idx < len(row_cells) else ''
            # 设置单元格字体
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=9)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    doc.add_paragraph()  # 表格后空一行
    return i, True


def markdown_to_docx(md_path, docx_path):
    doc = Document()

    # 设置默认中文字体
    style = doc.styles['Normal']
    style.font.name = '仿宋'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
    style.font.size = Pt(11)

    md_text = Path(md_path).read_text(encoding='utf-8')
    lines = md_text.splitlines()

    in_blockquote = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 块引用（保密提示）
        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            p = add_markdown_paragraph(doc, quote_text, size=9)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 水平分隔线
        if stripped == '---':
            doc.add_paragraph()
            i += 1
            continue

        # 表格
        if stripped.startswith('|'):
            next_idx, is_table = add_table_from_markdown(doc, lines, i)
            if is_table:
                i = next_idx
                continue

        # 标题
        heading_match = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            style_name = f'Heading {level}'
            p = doc.add_paragraph(style=style_name)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 6)
            p.paragraph_format.space_after = Pt(6)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, font_name='黑体', size=(18 if level == 1 else (14 if level == 2 else 12)), bold=True)
            i += 1
            continue

        # 普通段落
        add_markdown_paragraph(doc, stripped, size=11)
        i += 1

    # 页眉页脚或页边距可在此设置
    sections = doc.sections
    for section in sections:
        # 中国法院书面材料常用 A4；作为 standard_business_brief 的本案命名版式覆盖。
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(28)
        section.right_margin = Mm(25)

    doc.save(docx_path)
    print(f"已生成 Word 文档：{docx_path}")


if __name__ == '__main__':
    import argparse
    from docx.oxml import parse_xml
    parser = argparse.ArgumentParser(description='将 Markdown 转为 Word 文档')
    parser.add_argument('md_path', nargs='?', default=r'D:\aaa\张长松\md\质证意见-张长松对王书坡证据-2026-07-24.md')
    parser.add_argument('docx_path', nargs='?', default=r'D:\aaa\张长松\md\质证意见-张长松对王书坡证据-2026-07-24.docx')
    args = parser.parse_args()
    md_path = Path(args.md_path)
    docx_path = Path(args.docx_path)
    markdown_to_docx(md_path, docx_path)
